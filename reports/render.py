"""Выгрузки отчётов: XLSX, CSV (UTF-8 с BOM, «;»), PDF (ТЗ 9)."""
import csv
import io

from django.template.loader import render_to_string
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter


def header_lines(report, result, user):
    lines = [f'Отчёт {report.number}. {report.title}']
    lines += [f'{label}: {value}' for label, value in result.params + result.summary]
    lines.append(f'Сформирован: {timezone.localtime():%d.%m.%Y %H:%M}, {user.get_full_name() or user.get_username()}')
    return lines


def to_xlsx(report, result, user):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = f'Отчёт {report.number}'
    for line in header_lines(report, result, user):
        sheet.append([line])
    sheet['A1'].font = Font(bold=True, size=13)
    widths = {}
    for section in result.sections:
        sheet.append([])
        sheet.append([section.title])
        sheet.cell(sheet.max_row, 1).font = Font(bold=True)
        sheet.append(section.columns)
        for cell in sheet[sheet.max_row]:
            cell.font = Font(bold=True)
        for row in section.rows + ([section.totals] if section.totals else []):
            sheet.append(list(row))
            for index, value in enumerate(row, 1):
                text = str(value)
                widths[index] = max(widths.get(index, 8), min(60, max((len(x) for x in text.splitlines()), default=0)))
                if '\n' in text:
                    sheet.cell(sheet.max_row, index).alignment = Alignment(wrap_text=True, vertical='top')
        if section.totals:
            for cell in sheet[sheet.max_row]:
                cell.font = Font(bold=True)
        for index, title in enumerate(section.columns, 1):
            widths[index] = max(widths.get(index, 8), len(title) + 2)
    for index, width in widths.items():
        sheet.column_dimensions[get_column_letter(index)].width = width + 2
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def to_csv(report, result, user):
    buffer = io.StringIO()
    writer = csv.writer(buffer, delimiter=';', lineterminator='\r\n')
    for line in header_lines(report, result, user):
        writer.writerow([line])
    for section in result.sections:
        writer.writerow([])
        writer.writerow([section.title])
        writer.writerow(section.columns)
        writer.writerows(section.rows)
        if section.totals:
            writer.writerow(section.totals)
    return ('﻿' + buffer.getvalue()).encode('utf-8')


def to_pdf(report, result, user):
    from weasyprint import HTML
    html = render_to_string('reports/pdf.html', {
        'report': report, 'result': result, 'header': header_lines(report, result, user)})
    return HTML(string=html).write_pdf()
